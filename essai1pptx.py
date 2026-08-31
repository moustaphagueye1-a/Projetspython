from pptx import Presentation

# Chemin corrigé avec un seul jeu de guillemets
chemin_ppt = r"C:\Users\Admin\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\FC93334FB92A86E5ACBC7D69145290BC9A9BB455\transfers\2026-24\Cours Gestion d'entreprise agricole USSEIN 2025.pptx"

# Chargement de la présentation
ppt = Presentation(chemin_ppt)

# Ouverture (ou création) du fichier Markdown en mode écriture
with open("cours.md", "w", encoding="utf-8") as md:
    for i, slide in enumerate(ppt.slides, start=1):
        # Écrit le titre de la diapositive
        md.write(f"# Diapositive {i}\n\n")

        # Parcourt tous les éléments de la diapositive
        for shape in slide.shapes:
            # Si l'élément contient du texte et n'est pas vide
            if hasattr(shape, "text") and shape.text.strip():
                md.write(shape.text + "\n\n")

        # Ajoute une ligne de séparation entre les diapositives
        md.write("\n---\n\n")

print("Conversion terminée avec succès : cours.md")
